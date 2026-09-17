"""Word (.docx) export formatter for Radio & TV Story Segmenter.

Provides styled Microsoft Word document generation for transcripts with native OpenXML
comments/annotations compatible with LibreOffice Writer, Microsoft Word, and Google Docs:
- Arial typography and custom heading layouts
- Speaker labels with bold run formatting
- Timestamp annotations in muted styling
- Native OpenXML annotations/comments (word/comments.xml with commentRangeStart/End references)
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional
import xml.etree.ElementTree as ET
import zipfile

from prs_shared import format_time

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    try:
        from docx.enum.text import WD_COLOR_INDEX
    except ImportError:
        WD_COLOR_INDEX = None
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    WD_COLOR_INDEX = None


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _add_formatted_text_to_paragraph(
    p,
    text: str,
    words: Optional[List[Dict[str, Any]]] = None,
    block_fmt: Optional[Dict[str, Any]] = None,
    include_highlights: bool = True,
):
    """Applies word-level or block-level rich text formatting (bold, italic, underline, strike, highlight) to DOCX runs."""
    if words and len(words) > 0:
        for idx, w_item in enumerate(words):
            if isinstance(w_item, dict):
                w_str = w_item.get("word", "")
                b = bool(w_item.get("bold"))
                i = bool(w_item.get("italic"))
                u = bool(w_item.get("underline"))
                s = bool(w_item.get("strike"))
                hl = w_item.get("highlight") if include_highlights else None
            else:
                w_str = str(w_item)
                b = i = u = s = False
                hl = None

            if not w_str:
                continue

            prefix_space = " " if idx > 0 else ""
            r = p.add_run(prefix_space + w_str)
            if b:
                r.bold = True
            if i:
                r.italic = True
            if u:
                r.underline = True
            if s:
                try:
                    r.font.strike = True
                except Exception:
                    pass
            if hl and include_highlights:
                try:
                    if WD_COLOR_INDEX is not None:
                        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    else:
                        rPr = r._r.get_or_add_rPr()
                        rPr.append(ET.Element(f"{{{W_NS}}}highlight", {f"{{{W_NS}}}val": "yellow"}))
                except Exception:
                    pass
    else:
        r = p.add_run(text)
        if block_fmt:
            if block_fmt.get("bold"):
                r.bold = True
            if block_fmt.get("italic"):
                r.italic = True
            if block_fmt.get("underline"):
                r.underline = True
            if block_fmt.get("strike"):
                try:
                    r.font.strike = True
                except Exception:
                    pass
            if block_fmt.get("highlight") and include_highlights:
                try:
                    if WD_COLOR_INDEX is not None:
                        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    else:
                        rPr = r._r.get_or_add_rPr()
                        rPr.append(ET.Element(f"{{{W_NS}}}highlight", {f"{{{W_NS}}}val": "yellow"}))
                except Exception:
                    pass


def inject_native_comments_to_docx(
    docx_bytes: bytes,
    comments_map: Dict[int, List[Dict[str, Any]]],
    author: str = "Radio & TV Story Segmenter",
    initials: str = "RS"
) -> bytes:
    """Injects native OpenXML comment structures into a .docx ZIP archive.

    Ensures comments appear as standard native annotations in LibreOffice Writer,
    Microsoft Word, and Google Docs anchored directly to the selected text range.

    Args:
        docx_bytes: Raw bytes of the generated .docx file.
        comments_map: Mapping from paragraph 0-based body index to list of comment objects.
        author: Author label for the comment annotation.
        initials: Initials for the comment annotation.

    Returns:
        Modified .docx archive bytes with native comments.
    """
    if not comments_map:
        return docx_bytes

    in_buf = io.BytesIO(docx_bytes)
    out_buf = io.BytesIO()

    try:
        with zipfile.ZipFile(in_buf, "r") as in_zip, zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as out_zip:
            content_types_raw = in_zip.read("[Content_Types].xml").decode("utf-8")
            doc_rels_raw = in_zip.read("word/_rels/document.xml.rels").decode("utf-8")
            doc_xml_raw = in_zip.read("word/document.xml").decode("utf-8")

            # 1. Register comments.xml in [Content_Types].xml
            ET.register_namespace("", CONTENT_TYPES_NS)
            ct_tree = ET.fromstring(content_types_raw)
            overrides = ct_tree.findall(f"{{{CONTENT_TYPES_NS}}}Override")
            part_names = [ov.get("PartName") for ov in overrides]
            if "/word/comments.xml" not in part_names:
                ET.SubElement(ct_tree, f"{{{CONTENT_TYPES_NS}}}Override", {
                    "PartName": "/word/comments.xml",
                    "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
                })
            new_content_types_bytes = ET.tostring(ct_tree, encoding="utf-8", xml_declaration=True)

            # 2. Register relationship in word/_rels/document.xml.rels
            ET.register_namespace("", RELS_NS)
            rels_tree = ET.fromstring(doc_rels_raw)
            existing_rels = rels_tree.findall(f"{{{RELS_NS}}}Relationship")
            has_comments_rel = any(
                rel.get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
                for rel in existing_rels
            )
            if not has_comments_rel:
                max_num = 0
                for rel in existing_rels:
                    rid = rel.get("Id", "")
                    if rid.startswith("rId") and rid[3:].isdigit():
                        max_num = max(max_num, int(rid[3:]))
                new_rel_id = f"rId{max_num + 1}"
                ET.SubElement(rels_tree, f"{{{RELS_NS}}}Relationship", {
                    "Id": new_rel_id,
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
                    "Target": "comments.xml"
                })
            new_doc_rels_bytes = ET.tostring(rels_tree, encoding="utf-8", xml_declaration=True)

            # 3. Parse and annotate word/document.xml
            ET.register_namespace("w", W_NS)
            doc_tree = ET.fromstring(doc_xml_raw)
            body = doc_tree.find(f"{{{W_NS}}}body")
            comments_xml_entries: List[tuple[str, str, str, str]] = []

            if body is not None:
                paragraphs = body.findall(f"{{{W_NS}}}p")
                comment_id_counter = 0

                for p_idx, p in enumerate(paragraphs):
                    if p_idx in comments_map and comments_map[p_idx]:
                        for c_item in comments_map[p_idx]:
                            c_text = c_item.get("text", "").strip()
                            if not c_text:
                                continue
                            sel_text = c_item.get("selected_text", "").strip()
                            c_author = c_item.get("author") or author
                            cid = str(comment_id_counter)
                            comment_id_counter += 1

                            iso_now = c_item.get("date") or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
                            comments_xml_entries.append((cid, c_text, iso_now, c_author))

                            runs = p.findall(f"{{{W_NS}}}r")
                            if not runs:
                                continue

                            # Try anchoring to exact selected_text if present in paragraph
                            anchored = False
                            if sel_text:
                                run_info = []
                                current_offset = 0
                                for r in runs:
                                    t_elems = r.findall(f"{{{W_NS}}}t")
                                    r_text = "".join(t.text or "" for t in t_elems)
                                    run_info.append((r, r_text, current_offset))
                                    current_offset += len(r_text)

                                full_p_text = "".join(ri[1] for ri in run_info)
                                match_start = full_p_text.find(sel_text)
                                if match_start >= 0:
                                    match_end = match_start + len(sel_text)
                                    # Split runs at match boundaries
                                    p_children = list(p)
                                    for r, r_text, r_start in run_info:
                                        r_end = r_start + len(r_text)
                                        if not r_text:
                                            continue
                                        splits = []
                                        if r_start < match_start < r_end:
                                            splits.append(match_start - r_start)
                                        if r_start < match_end < r_end:
                                            splits.append(match_end - r_start)

                                        if splits:
                                            splits = sorted(list(set(splits)))
                                            idx_in_p = p_children.index(r)
                                            p.remove(r)

                                            last_idx = 0
                                            r_pr = r.find(f"{{{W_NS}}}rPr")
                                            for s_idx in splits + [len(r_text)]:
                                                chunk = r_text[last_idx:s_idx]
                                                if chunk:
                                                    new_r = ET.Element(f"{{{W_NS}}}r")
                                                    if r_pr is not None:
                                                        new_r.append(ET.fromstring(ET.tostring(r_pr)))
                                                    new_t = ET.SubElement(new_r, f"{{{W_NS}}}t", {f"{{{W_NS}}}space": "preserve"})
                                                    new_t.text = chunk
                                                    p.insert(idx_in_p, new_r)
                                                    idx_in_p += 1
                                                last_idx = s_idx
                                            p_children = list(p)

                                    # Re-query runs after splits
                                    new_run_info = []
                                    curr_off = 0
                                    for r in p.findall(f"{{{W_NS}}}r"):
                                        t_elems = r.findall(f"{{{W_NS}}}t")
                                        r_text = "".join(t.text or "" for t in t_elems)
                                        new_run_info.append((r, r_text, curr_off))
                                        curr_off += len(r_text)

                                    target_runs = [r for r, r_text, r_start in new_run_info if r_text and (r_start >= match_start and r_start + len(r_text) <= match_end)]
                                    if target_runs:
                                        first_target = target_runs[0]
                                        last_target = target_runs[-1]
                                        p_children = list(p)
                                        if first_target in p_children and last_target in p_children:
                                            start_idx = p_children.index(first_target)
                                            start_elem = ET.Element(f"{{{W_NS}}}commentRangeStart", {f"{{{W_NS}}}id": cid})
                                            end_elem = ET.Element(f"{{{W_NS}}}commentRangeEnd", {f"{{{W_NS}}}id": cid})

                                            ref_run = ET.Element(f"{{{W_NS}}}r")
                                            rpr = ET.SubElement(ref_run, f"{{{W_NS}}}rPr")
                                            ET.SubElement(rpr, f"{{{W_NS}}}rStyle", {f"{{{W_NS}}}val": "CommentReference"})
                                            ET.SubElement(ref_run, f"{{{W_NS}}}commentReference", {f"{{{W_NS}}}id": cid})

                                            p.insert(start_idx, start_elem)
                                            p_children_after_start = list(p)
                                            end_idx = p_children_after_start.index(last_target) + 1
                                            p.insert(end_idx, end_elem)
                                            p.insert(end_idx + 1, ref_run)
                                            anchored = True

                            if not anchored:
                                body_runs = [
                                    r for r in runs
                                    if not (r.find(f".//{{{W_NS}}}color") is not None and "787878" in (r.find(f".//{{{W_NS}}}color").get(f"{{{W_NS}}}val", "")))
                                ]
                                if not body_runs:
                                    body_runs = runs

                                first_target = body_runs[0]
                                last_target = body_runs[-1]
                                p_children = list(p)

                                if first_target in p_children and last_target in p_children:
                                    start_idx = p_children.index(first_target)

                                    start_elem = ET.Element(f"{{{W_NS}}}commentRangeStart", {f"{{{W_NS}}}id": cid})
                                    end_elem = ET.Element(f"{{{W_NS}}}commentRangeEnd", {f"{{{W_NS}}}id": cid})

                                    ref_run = ET.Element(f"{{{W_NS}}}r")
                                    rpr = ET.SubElement(ref_run, f"{{{W_NS}}}rPr")
                                    ET.SubElement(rpr, f"{{{W_NS}}}rStyle", {f"{{{W_NS}}}val": "CommentReference"})
                                    ET.SubElement(ref_run, f"{{{W_NS}}}commentReference", {f"{{{W_NS}}}id": cid})

                                    p.insert(start_idx, start_elem)
                                    p_children_after_start = list(p)
                                    end_idx = p_children_after_start.index(last_target) + 1
                                    p.insert(end_idx, end_elem)
                                    p.insert(end_idx + 1, ref_run)

            # 4. Generate word/comments.xml if any comments exist
            comments_xml_bytes = b""
            if comments_xml_entries:
                comments_root = ET.Element(f"{{{W_NS}}}comments")
                for cid, c_text, iso_now, c_auth in comments_xml_entries:
                    c_elem = ET.SubElement(comments_root, f"{{{W_NS}}}comment", {
                        f"{{{W_NS}}}id": cid,
                        f"{{{W_NS}}}author": c_auth,
                        f"{{{W_NS}}}initials": initials,
                        f"{{{W_NS}}}date": iso_now,
                    })
                    cp = ET.SubElement(c_elem, f"{{{W_NS}}}p")
                    cpPr = ET.SubElement(cp, f"{{{W_NS}}}pPr")
                    ET.SubElement(cpPr, f"{{{W_NS}}}pStyle", {f"{{{W_NS}}}val": "CommentText"})

                    cr_ref = ET.SubElement(cp, f"{{{W_NS}}}r")
                    cr_ref_pr = ET.SubElement(cr_ref, f"{{{W_NS}}}rPr")
                    ET.SubElement(cr_ref_pr, f"{{{W_NS}}}rStyle", {f"{{{W_NS}}}val": "CommentReference"})
                    ET.SubElement(cr_ref, f"{{{W_NS}}}annotationRef")

                    cr_text = ET.SubElement(cp, f"{{{W_NS}}}r")
                    ct = ET.SubElement(cr_text, f"{{{W_NS}}}t")
                    ct.text = c_text

                comments_xml_bytes = ET.tostring(comments_root, encoding="utf-8", xml_declaration=True)

            new_doc_xml_bytes = ET.tostring(doc_tree, encoding="utf-8", xml_declaration=True)

            # Copy all files into output zip
            for item in in_zip.infolist():
                if item.filename == "[Content_Types].xml":
                    out_zip.writestr(item.filename, new_content_types_bytes)
                elif item.filename == "word/_rels/document.xml.rels":
                    out_zip.writestr(item.filename, new_doc_rels_bytes)
                elif item.filename == "word/document.xml":
                    out_zip.writestr(item.filename, new_doc_xml_bytes)
                else:
                    out_zip.writestr(item.filename, in_zip.read(item.filename))

            if comments_xml_entries and comments_xml_bytes:
                out_zip.writestr("word/comments.xml", comments_xml_bytes)

        return out_buf.getvalue()
    except Exception as exc:
        print(f"[DOCX_EXPORT] Warning: Failed to inject native comments: {exc}. Using base DOCX output.")
        return docx_bytes


def create_story_docx(
    title: str,
    blocks: List[Dict[str, Any]],
    output_path: Path | str,
    media_name: str = "",
    start_time: float = 0.0,
    end_time: float = 0.0,
    include_speakers: bool = True,
    include_timestamps: bool = True,
    include_comments: bool = True,
    include_highlights: bool = True,
    lang_code: str = "en",
    source_segments: Optional[List[Dict[str, Any]]] = None,
) -> bool:
    """Creates a formatted Word (.docx) document with native OpenXML comments from story blocks and metadata."""
    if not DOCX_AVAILABLE:
        return False

    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)

    lang_label = " (Spanish)" if lang_code == "es" else (" (English)" if lang_code == "en" else "")
    document.add_heading(f"{title}{lang_label}", 0)

    # Paragraph index tracking for native OpenXML comments
    # 0 = heading paragraph
    body_para_idx = 1
    if media_name:
        time_range = f" ({format_time(start_time, False)} - {format_time(end_time, False)})" if end_time > 0 else ""
        document.add_paragraph(f"Recording: {media_name}{time_range}")
        body_para_idx += 1

    comments_map: Dict[int, List[Dict[str, Any]]] = {}
    last_speaker = None
    processed_seg_comments = set()

    for block in blocks:
        speaker = (block.get("speaker") or "").strip() if include_speakers else ""
        p_text = block.get("text", "").strip()
        if not p_text:
            continue

        p = document.add_paragraph()
        if include_timestamps and "start" in block and block["start"] is not None:
            r_time = p.add_run(f"[{format_time(block['start'], False)}] ")
            r_time.font.color.rgb = RGBColor(120, 120, 120)

        is_speaker_change = block.get("is_speaker_change", (speaker != last_speaker))
        if speaker and is_speaker_change and speaker != last_speaker:
            r_spk = p.add_run(f"{speaker}: ")
            r_spk.bold = True
            last_speaker = speaker

        words = block.get("words")
        if not words and source_segments:
            src_indices = block.get("source_indices")
            if src_indices is None and block.get("_source_index") is not None:
                src_indices = [block["_source_index"]]
            if src_indices:
                words = []
                for s_idx in src_indices:
                    if 0 <= s_idx < len(source_segments):
                        seg_words = source_segments[s_idx].get("words")
                        if seg_words and isinstance(seg_words, list):
                            words.extend(seg_words)

        _add_formatted_text_to_paragraph(p, p_text, words=words, block_fmt=block, include_highlights=include_highlights)
        p.paragraph_format.space_after = Pt(6)

        if include_comments:
            c_items = []
            # 1. Inspect block comments_data or comments
            if block.get("comments_data"):
                cd = block["comments_data"]
                if isinstance(cd, list):
                    c_items.extend(cd)
                elif isinstance(cd, dict):
                    c_items.append(cd)

            # 2. Inspect source_indices
            src_indices = block.get("source_indices")
            if src_indices is None and block.get("_source_index") is not None:
                src_indices = [block["_source_index"]]

            if src_indices and source_segments:
                for s_idx in src_indices:
                    if 0 <= s_idx < len(source_segments):
                        seg = source_segments[s_idx]
                        c_text = (seg.get("comments") or seg.get("notes", "")).strip()
                        if c_text:
                            processed_seg_comments.add((s_idx, c_text))
                            sel_q = (seg.get("comment_selected_text") or "").strip()
                            if not any(ci.get("text") == c_text and ci.get("selected_text") == sel_q for ci in c_items):
                                c_items.append({
                                    "text": c_text,
                                    "selected_text": sel_q,
                                })

            # 3. Fallback to block.get("comments")
            if not c_items:
                raw_c = str(block.get("comments") or block.get("notes", "")).strip()
                if raw_c:
                    c_items.append({
                        "text": raw_c,
                        "selected_text": str(block.get("comment_selected_text") or "").strip()
                    })

            if c_items:
                comments_map[body_para_idx] = c_items

        body_para_idx += 1

    # Ensure any remaining unprocessed comments in source_segments are exported
    if include_comments and source_segments:
        unprocessed = []
        for idx, seg in enumerate(source_segments):
            c_text = (seg.get("comments") or seg.get("notes", "")).strip()
            if c_text and (idx, c_text) not in processed_seg_comments:
                unprocessed.append((idx, seg, c_text))

        if unprocessed:
            for idx, seg, c_text in unprocessed:
                p = document.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(f"[{format_time(seg.get('start', 0.0))}] {seg.get('speaker', '')}: {seg.get('text', '')}")
                r.font.size = Pt(10)
                comments_map[body_para_idx] = [{
                    "text": c_text,
                    "selected_text": (seg.get("comment_selected_text") or "").strip(),
                }]
                body_para_idx += 1

    # Save to memory buffer
    mem_buf = io.BytesIO()
    document.save(mem_buf)
    docx_bytes = mem_buf.getvalue()

    # If comments exist, inject native OpenXML annotations
    if include_comments and comments_map:
        docx_bytes = inject_native_comments_to_docx(docx_bytes, comments_map)

    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    out_file.write_bytes(docx_bytes)
    return True

